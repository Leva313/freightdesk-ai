"""Generate the binary sample requests (PDF, spreadsheet, chat screenshot).

Run with the dev dependencies installed:  python samples/make_samples.py
"""

from datetime import datetime, timezone
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent / "inputs"


def transport_order_pdf() -> None:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.set_creation_date(datetime(2026, 10, 12, 9, 0, tzinfo=timezone.utc))
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(15, 40, 80)
    pdf.cell(0, 10, "Van Dijk Logistics B.V.", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 5, "Waalhaven Oostzijde 85, 3087 BM Rotterdam, The Netherlands | planning@vandijk-logistics.nl | +31 10 555 2290",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 9, "TRANSPORT ORDER  No. VDL-77310", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, "Date: 12.10.2026   Contact: Pieter de Boer, Planning", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    def section(title: str, rows: list[tuple[str, str]]) -> None:
        pdf.set_fill_color(230, 236, 245)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 8, "  " + title, fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        for key, value in rows:
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(48, 7, key)
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 7, value, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

    section("1. Loading", [
        ("Shipper", "Van Dijk Logistics B.V. - Warehouse 3"),
        ("Address", "Waalhaven Oostzijde 85, 3087 BM Rotterdam, NL"),
        ("Date / time", "20.10.2026, 09:00 - 15:00"),
        ("Loading reference", "VDL-77310-L"),
    ])
    section("2. Unloading", [
        ("Consignee", "Kaufland Ceska republika v.o.s. - DC Praha"),
        ("Address", "Belohorska 2428/203, 169 00 Praha 6, CZ"),
        ("Date / time", "22.10.2026, 06:00 - 12:00 (slot booking required)"),
        ("Unloading reference", "KCZ-PO-4481920"),
    ])
    section("3. Goods", [
        ("Description", "Consumer electronics (TV sets, boxed)"),
        ("Packaging", "30 EUR pallets, not stackable"),
        ("Gross weight", "11,250 kg"),
        ("Loading meters", "13.6 ldm (full truck load)"),
        ("Dangerous goods", "No"),
    ])
    section("4. Vehicle and price", [
        ("Vehicle", "Mega trailer 100 m3, curtainsider, 2 x 24 straps"),
        ("Freight rate", "EUR 1,780.00 all-in, payment 30 days"),
        ("Instructions", "Driver to call 1 hour before arrival. High value goods - no unattended parking."),
    ])
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(110, 110, 110)
    pdf.multi_cell(0, 4, "Carriage under CMR convention. Please confirm this order by return email with truck and trailer plates.")
    pdf.output(OUT / "05_transport_order_pdf_rotterdam_prague.pdf")


def booking_sheet_xlsx() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Booking"
    rows = [
        ("FREIGHT BOOKING FORM", ""),
        ("Customer", "Agrícola Mediterránea S.L."),
        ("Contact", "Laura Martínez"),
        ("Email", "l.martinez@agromed.es"),
        ("Phone", "+34 93 555 7781"),
        ("Order ref", "AGM-5521"),
        ("", ""),
        ("Loading place", "Mercabarna, Carrer Longitudinal 6, 08040 Barcelona, ES"),
        ("Loading date", "13/10/2026"),
        ("Loading time", "14:00-20:00"),
        ("Unloading place", "Großmarkthalle München, Schäftlarnstraße 10, 81371 München, DE"),
        ("Unloading date", "15/10/2026"),
        ("Unloading time", "04:00-08:00"),
        ("", ""),
        ("Goods", "Fresh citrus fruit (oranges, mandarins)"),
        ("Pallets", 33),
        ("Pallet type", "EUR"),
        ("Gross weight (kg)", 21500),
        ("Temperature", "+4 / +8 °C"),
        ("Vehicle", "Reefer 13.6 m"),
        ("Freight rate", "2650 EUR"),
    ]
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="2F855A")
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 62
    for cell in sheet["A"][1:]:
        cell.font = Font(bold=True)
    workbook.save(OUT / "06_booking_sheet_xlsx_barcelona_munich.xlsx")


def chat_screenshot_png() -> None:
    width, height = 720, 760
    image = Image.new("RGB", (width, height), "#ECE5DD")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(font_path, 24)
    small = ImageFont.truetype(font_path, 18)
    bold = ImageFont.truetype(bold_path, 26)

    draw.rectangle((0, 0, width, 90), fill="#075E54")
    draw.ellipse((24, 18, 78, 72), fill="#cfd8dc")
    draw.text((96, 18), "Tomasz - AutoParts Trading", font=bold, fill="white")
    draw.text((96, 54), "online", font=small, fill="#d0f0e8")

    message = [
        "Hi! Need a truck from Antwerp port to Warsaw",
        "",
        "Pickup: Thursday 15.10, 10:00-16:00",
        "DP World Antwerp Gateway, Scheldelaan 450,",
        "2040 Antwerpen",
        "",
        "Delivery: AutoParts Trading Sp. z o.o.,",
        "ul. Marywilska 34, 03-228 Warszawa",
        "Delivery Monday 19.10 morning",
        "",
        "18 pallets auto parts (brake discs), about 14 t",
        "Tautliner is fine. Container ref MSKU 4471932",
        "",
        "My number +32 470 55 12 34",
    ]
    top, left, line = 130, 40, 34
    bubble_height = len(message) * line + 50
    draw.rounded_rectangle((left - 16, top - 20, width - 70, top + bubble_height), radius=18, fill="white")
    for i, text in enumerate(message):
        draw.text((left, top + i * line), text, font=font, fill="#111")
    draw.text((width - 160, top + bubble_height - 34), "09:42", font=small, fill="#888")
    image.save(OUT / "07_chat_screenshot_png_antwerp_warsaw.png", optimize=True)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    transport_order_pdf()
    booking_sheet_xlsx()
    chat_screenshot_png()
    print("Samples written to", OUT)
