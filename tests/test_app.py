import io
import json
from datetime import date
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, documents, extractor
from app.schema import FreightOrder, missing_fields

SAMPLES = extractor.SAMPLES_DIR / "inputs"


@pytest.fixture(autouse=True)
def isolated_data(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DATA_DIR", tmp_path)
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(db, "UPLOADS_DIR", tmp_path / "uploads")
    monkeypatch.setenv("FREIGHTDESK_DEMO", "1")


@pytest.fixture
def client():
    from app.main import app

    return TestClient(app)


def load_expected(stem: str) -> FreightOrder:
    return FreightOrder.model_validate_json((extractor.SAMPLES_DIR / "expected" / f"{stem}.json").read_text())


def test_every_sample_has_a_valid_expected_result():
    for path in SAMPLES.iterdir():
        load_expected(path.stem)


@pytest.mark.parametrize("path", sorted(SAMPLES.iterdir()), ids=lambda p: p.suffix)
def test_documents_convert_each_sample_format(path):
    document = documents.from_upload(path.name, path.read_bytes())
    assert document.blocks
    if path.suffix == ".xlsx":
        assert "Mercabarna" in document.preview
    if path.suffix == ".pdf":
        assert document.blocks[0]["type"] == "document"


def test_unsupported_upload_is_rejected():
    with pytest.raises(documents.UnsupportedDocument):
        documents.from_upload("notes.docx", b"data")


def test_missing_fields_follow_adr_and_temperature_rules():
    order = load_expected("03_email_pl_adr_poznan_milano")
    assert missing_fields(order) == []
    order.cargo.adr_class = None
    order.cargo.un_number = None
    assert "cargo.adr_class" in missing_fields(order)

    reefer = load_expected("02_email_de_kuehltransport_hamburg_wien")
    reefer.cargo.temp_min_c = reefer.cargo.temp_max_c = None
    assert "cargo.temp_min_c" in missing_fields(reefer)


def test_demo_mode_rejects_unknown_documents():
    with pytest.raises(extractor.DemoModeError):
        extractor.extract(documents.from_text("truck please"), b"truck please")


def test_sample_flow_status_edit_confirm_and_export(client):
    response = client.post("/samples/04_email_ru_riga_warsaw_missing_info.txt", follow_redirects=False)
    assert response.status_code == 303
    order_id = int(response.headers["location"].rsplit("/", 1)[1])

    order = db.get_order(order_id)
    assert order.status == "needs_info"
    assert order.data.clarification_email
    assert client.get(f"/orders/{order_id}").status_code == 200

    response = client.post(f"/orders/{order_id}", data={"action": "confirm", "cargo.description": "Boxed furniture"})
    assert "Fill in before confirming" in response.text
    assert db.get_order(order_id).status == "needs_info"

    response = client.post(
        f"/orders/{order_id}",
        data={"action": "confirm", "pickup.date_from": "2026-10-20", "cargo.weight_kg": "8200"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    confirmed = db.get_order(order_id)
    assert confirmed.status == "confirmed"
    assert confirmed.data.cargo.weight_kg == 8200
    assert confirmed.data.cargo.stackable is True  # fields that were not posted keep their values

    export = client.get("/export.xlsx")
    sheet = load_workbook(io.BytesIO(export.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1][0] == order_id and rows[1][17] == 8200


def test_image_sample_is_stored_and_served(client):
    response = client.post("/samples/07_chat_screenshot_png_antwerp_warsaw.png")
    assert response.status_code == 200
    order = db.list_orders()[0]
    assert order.source_kind == "image" and order.status == "new"
    assert client.get(f"/orders/{order.id}/source").headers["content-type"] == "image/png"


def test_api_lists_orders(client):
    client.post("/samples/01_email_en_paper_manchester_lyon.txt")
    data = client.get("/api/orders").json()
    assert data[0]["order"]["pickup"]["city"] == "Manchester"
    assert json.dumps(data)


def test_live_extraction_sends_document_and_parses_result(monkeypatch):
    expected = load_expected("05_transport_order_pdf_rotterdam_prague")
    calls = {}

    class FakeMessages:
        def parse(self, **kwargs):
            calls.update(kwargs)
            return SimpleNamespace(stop_reason="end_turn", parsed_output=expected)

    monkeypatch.setattr(
        extractor.anthropic, "Anthropic", lambda: SimpleNamespace(beta=SimpleNamespace(messages=FakeMessages()))
    )
    pdf = SAMPLES / "05_transport_order_pdf_rotterdam_prague.pdf"
    result = extractor._claude_extract(documents.from_upload(pdf.name, pdf.read_bytes()), received=date(2026, 10, 12))

    assert result == expected
    assert calls["model"] == extractor.MODEL
    assert calls["output_format"] is FreightOrder
    assert calls["messages"][0]["content"][0]["type"] == "document"
    assert "2026-10-12" in calls["messages"][0]["content"][-1]["text"]


def test_hosted_demo_seeds_on_start_and_never_calls_the_api(monkeypatch):
    from app.main import app

    monkeypatch.setenv("FREIGHTDESK_SEED_ON_START", "1")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-should-not-be-used")
    assert not extractor.is_live()

    with TestClient(app) as client:  # runs the lifespan startup
        assert len(db.list_orders()) == len(list(SAMPLES.iterdir()))
        assert "Live demo." in client.get("/").text
        response = client.post("/orders", data={"text": "Need a truck Paris to Madrid tomorrow"})
        assert "uploading your own documents is disabled" in response.text
    assert len(db.list_orders()) == len(list(SAMPLES.iterdir()))
