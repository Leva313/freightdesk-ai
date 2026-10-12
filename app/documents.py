"""Turn an uploaded file or pasted text into Claude content blocks plus a text preview for the UI."""

from __future__ import annotations

import base64
import email
import io
from dataclasses import dataclass
from email import policy
from pathlib import PurePath

from openpyxl import load_workbook

IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
TEXT_SUFFIXES = {".txt", ".csv", ".md"}
ACCEPTED_SUFFIXES = sorted({*IMAGE_TYPES, *TEXT_SUFFIXES, ".pdf", ".eml", ".xlsx"})
MAX_UPLOAD_BYTES = 20 * 1024 * 1024


class UnsupportedDocument(ValueError):
    pass


@dataclass
class Document:
    kind: str  # text | email | pdf | image | excel
    blocks: list[dict]
    preview: str | None  # text shown next to the extracted fields; None for PDF and images


def from_text(text: str) -> Document:
    text = text.strip()
    if not text:
        raise UnsupportedDocument("The request is empty.")
    return Document("text", [{"type": "text", "text": text}], text)


def from_upload(filename: str, data: bytes) -> Document:
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnsupportedDocument("File is larger than 20 MB.")
    suffix = PurePath(filename).suffix.lower()

    if suffix == ".pdf":
        block = {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": _b64(data)},
        }
        return Document("pdf", [block], None)

    if suffix in IMAGE_TYPES:
        block = {
            "type": "image",
            "source": {"type": "base64", "media_type": IMAGE_TYPES[suffix], "data": _b64(data)},
        }
        return Document("image", [block], None)

    if suffix == ".eml":
        text = _email_to_text(data)
        return Document("email", [{"type": "text", "text": text}], text)

    if suffix == ".xlsx":
        text = _workbook_to_text(data)
        return Document("excel", [{"type": "text", "text": text}], text)

    if suffix in TEXT_SUFFIXES:
        return from_text(data.decode("utf-8", errors="replace"))

    raise UnsupportedDocument(f"Unsupported file type '{suffix or filename}'. Use: {', '.join(ACCEPTED_SUFFIXES)}.")


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode("ascii")


def _email_to_text(data: bytes) -> str:
    message = email.message_from_bytes(data, policy=policy.default)
    body = message.get_body(preferencelist=("plain", "html"))
    content = body.get_content() if body else ""
    headers = [f"{name}: {message[name]}" for name in ("From", "To", "Date", "Subject") if message[name]]
    return "\n".join(headers) + "\n\n" + content.strip()


def _workbook_to_text(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cells):
                rows.append("\t".join(cells).rstrip())
        if rows:
            parts.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
    if not parts:
        raise UnsupportedDocument("The spreadsheet has no data.")
    return "\n\n".join(parts)
