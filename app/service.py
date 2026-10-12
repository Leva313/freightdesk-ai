"""Glue between documents, the extractor and storage."""

from __future__ import annotations

import io
import uuid
from pathlib import PurePath

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import db, documents, extractor, routing, schedule
from .db import Order
from .schema import FreightOrder


def ingest_upload(filename: str, data: bytes) -> int:
    document = documents.from_upload(filename, data)
    order = extractor.extract(document, data)
    stored = _store_upload(filename, data) if document.kind in ("pdf", "image") else None
    order_id = db.create_order(order, filename, document.kind, stored, document.preview, demo=not extractor.is_live())
    refresh_route(order_id)
    return order_id


def ingest_text(text: str) -> int:
    document = documents.from_text(text)
    raw = text.strip().encode("utf-8")
    order = extractor.extract(document, raw)
    first_line = text.strip().splitlines()[0][:60]
    order_id = db.create_order(order, first_line, document.kind, None, document.preview, demo=not extractor.is_live())
    refresh_route(order_id)
    return order_id


ROUTE_FIELDS = ("address", "postal_code", "city", "country")


def route_inputs_changed(before: FreightOrder, after: FreightOrder) -> bool:
    return any(
        getattr(getattr(before, stop), f) != getattr(getattr(after, stop), f)
        for stop in ("pickup", "delivery")
        for f in ROUTE_FIELDS
    ) or (before.cargo.adr, before.cargo.adr_class) != (after.cargo.adr, after.cargo.adr_class)


def refresh_route(order_id: int) -> None:
    """Fetch the truck route for an order and store it, or store why it could not be fetched."""
    order = db.get_order(order_id)
    if order is None:
        return
    try:
        route = routing.route_for(order.data, order.source_name, schedule.loading_start(order.data))
    except routing.RoutingError as exc:
        db.set_route(order_id, None, str(exc))
        return
    db.set_route(order_id, route, None if route or routing.has_key() else "Add HERE_API_KEY to .env to plan routes.")


def _store_upload(filename: str, data: bytes) -> str:
    db.UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    name = f"{uuid.uuid4().hex}{PurePath(filename).suffix.lower()}"
    (db.UPLOADS_DIR / name).write_bytes(data)
    return name


EXPORT_COLUMNS = [
    ("Order ID", lambda o: o.id),
    ("Customer", lambda o: o.data.customer.company),
    ("Customer ref", lambda o: o.data.customer_reference),
    ("Pickup company", lambda o: o.data.pickup.company),
    ("Pickup address", lambda o: o.data.pickup.address),
    ("Pickup postal code", lambda o: o.data.pickup.postal_code),
    ("Pickup city", lambda o: o.data.pickup.city),
    ("Pickup country", lambda o: o.data.pickup.country),
    ("Pickup date", lambda o: o.data.pickup.date_from),
    ("Pickup window", lambda o: o.data.pickup.time_window),
    ("Delivery company", lambda o: o.data.delivery.company),
    ("Delivery address", lambda o: o.data.delivery.address),
    ("Delivery postal code", lambda o: o.data.delivery.postal_code),
    ("Delivery city", lambda o: o.data.delivery.city),
    ("Delivery country", lambda o: o.data.delivery.country),
    ("Delivery date", lambda o: o.data.delivery.date_from),
    ("Cargo", lambda o: o.data.cargo.description),
    ("Weight kg", lambda o: o.data.cargo.weight_kg),
    ("Pallets", lambda o: o.data.cargo.pallets),
    ("Pallet type", lambda o: o.data.cargo.pallet_type),
    ("LDM", lambda o: o.data.cargo.loading_meters),
    ("ADR", lambda o: "yes" if o.data.cargo.adr else "no"),
    ("ADR class / UN", lambda o: " / ".join(p for p in (o.data.cargo.adr_class, o.data.cargo.un_number) if p)),
    ("Temp °C", lambda o: _temp_range(o)),
    ("Vehicle", lambda o: o.data.vehicle_type),
    ("Distance km", lambda o: o.route.distance_km if o.route else None),
    ("Driving h", lambda o: round(o.route.drive_seconds / 3600, 1) if o.route else None),
    ("ETA", lambda o: _eta(o)),
    ("On time", lambda o: (trip_plan(o).status if o.route else None)),
    ("Price", lambda o: o.data.price),
    ("Currency", lambda o: o.data.currency),
    ("Notes", lambda o: o.data.notes),
]


def _eta(order: Order) -> str | None:
    plan = trip_plan(order)
    return plan.arrival.strftime("%Y-%m-%d %H:%M") if plan and plan.arrival else None


def _temp_range(order: Order) -> str | None:
    cargo = order.data.cargo
    if not cargo.temperature_controlled:
        return None
    low = "" if cargo.temp_min_c is None else f"{cargo.temp_min_c:g}"
    high = "" if cargo.temp_max_c is None else f"{cargo.temp_max_c:g}"
    return f"{low}..{high}"


def trip_plan(order: Order) -> schedule.TripPlan | None:
    return schedule.plan_trip(order.data, order.route.drive_seconds) if order.route else None


def export_xlsx(orders: list[Order]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append([name for name, _ in EXPORT_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")
    for order in orders:
        sheet.append([getter(order) for _, getter in EXPORT_COLUMNS])
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 40)
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
